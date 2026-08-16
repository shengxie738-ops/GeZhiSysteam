import asyncio
import io
import json
import os
import unittest
from unittest.mock import AsyncMock, MagicMock, patch
from urllib.parse import unquote, quote

import docx
from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

os.environ.setdefault("RAGFLOW_API_KEY", "test")
os.environ.setdefault("RAGFLOW_BASE_URL", "http://localhost")
os.environ.setdefault("RAGFLOW_AGENT_ID", "test")
os.environ.setdefault("RAGFLOW_CHAT_ID", "test")
os.environ.setdefault("RAGFLOW_DATASET_ID", "test")
os.environ.setdefault("RAGFLOW_PUBLIC_DATASET_IDS", "")
os.environ.setdefault("OPENAI_API_KEY", "test")
os.environ.setdefault("OPENAI_API_BASE", "http://localhost")
os.environ.setdefault("AI_LESSON_PREP_API_KEY", "lesson-test-key")
os.environ.setdefault("AI_LESSON_PREP_BASE_URL", "https://example.test/chat/completions")
os.environ.setdefault("AI_LESSON_PREP_MODEL", "glm-4.5-air")

from app.api.endpoints.teacher_lesson_prep import (
    generate_lesson_plan,
    get_lesson_prep_config,
    get_lesson_prep_draft,
    list_lesson_prep_drafts,
    list_lesson_prep_resources,
    router,
    save_lesson_prep_draft,
    search_lesson_prep_resources,
    summarize_lesson_prep_resources,
)
from app.core.config import settings
from app.core.security import create_access_token
from app.models.domain_record import DomainRecord
from app.schemas.teacher_lesson_prep import (
    AILessonPlanResponse,
    AILessonSummaryResponse,
    LessonPlanExportRequest,
    LessonPlanGenerateRequest,
    LessonPrepDraftRequest,
    LessonPrepSearchRequest,
    LessonPrepSummaryRequest,
    TeachingStageExport,
)
from app.services.teacher_lesson_prep.ai_client import LessonPrepAIClient
from app.services.teacher_lesson_prep.service import TeacherLessonPrepService


class FakeResponse:
    def __init__(self, payload):
        self.payload = payload

    def raise_for_status(self):
        return None

    def json(self):
        return self.payload


class FakeAsyncClient:
    last_request = None
    response_payload = None

    def __init__(self, **kwargs):
        self.kwargs = kwargs

    async def __aenter__(self):
        return self

    async def __aexit__(self, exc_type, exc, tb):
        return None

    async def post(self, url, **kwargs):
        FakeAsyncClient.last_request = {"url": url, **kwargs, "client": self.kwargs}
        return FakeResponse(FakeAsyncClient.response_payload)


class TeacherLessonPrepApiTest(unittest.TestCase):
    def setUp(self):
        engine = create_engine("sqlite:///:memory:")
        DomainRecord.__table__.create(bind=engine)
        self.db = sessionmaker(bind=engine)()
        self.teacher_auth = f"Bearer {create_access_token('teacher-1', 'teacher')}"
        self.student_auth = f"Bearer {create_access_token('student-1', 'student')}"

    def tearDown(self):
        self.db.close()

    def test_student_cannot_access_teacher_lesson_prep(self):
        with self.assertRaises(HTTPException) as ctx:
            asyncio.run(get_lesson_prep_config(self.student_auth))
        self.assertEqual(ctx.exception.status_code, 403)

    def test_config_exposes_limits_but_not_credentials(self):
        result = asyncio.run(get_lesson_prep_config(self.teacher_auth))

        self.assertEqual(result["data"]["model"], "glm-4.5-air")
        self.assertTrue(result["data"]["ai_ready"])
        self.assertEqual(result["data"]["max_input_tokens"], 81920)
        self.assertEqual(result["data"]["max_output_tokens"], 49152)
        serialized = str(result).lower()
        self.assertNotIn("api_key", serialized)
        self.assertNotIn("lesson-test-key", serialized)
        self.assertNotIn("base_url", serialized)

    def test_catalog_and_search_delegate_to_catalog_service(self):
        fake_catalog = {
            "summary": {"total": 97, "pdf": 58, "slides": 39, "searchable": 58},
            "resources": [{"id": "courseware-1", "name": "linear-list.pdf"}],
        }
        fake_search = {
            "query": "linear list",
            "matches": [{"resource_id": "courseware-1", "page": 6, "excerpt": "linear list"}],
        }
        with patch("app.api.endpoints.teacher_lesson_prep.lesson_prep_service.list_resources", return_value=fake_catalog), patch(
            "app.api.endpoints.teacher_lesson_prep.lesson_prep_service.search", return_value=fake_search
        ):
            catalog_result = asyncio.run(list_lesson_prep_resources(self.teacher_auth, None, None, None))
            search_result = asyncio.run(
                search_lesson_prep_resources(
                    LessonPrepSearchRequest(query="linear list", resource_ids=["courseware-1"]),
                    self.teacher_auth,
                )
            )

        self.assertEqual(catalog_result["data"]["summary"]["total"], 97)
        self.assertEqual(search_result["data"]["matches"][0]["page"], 6)

    def test_summary_and_plan_generation_use_ai_without_exposing_secret(self):
        summary_payload = {
            "content": "课件围绕线性表定义、顺序表插入和链表删除展开。",
            "citations": [{"resource_id": "courseware-1", "page": 6, "excerpt": "线性表定义"}],
            "model": "glm-4.5-air",
        }
        plan_payload = {
            "title": "线性表教学设计",
            "duration_minutes": 45,
            "objectives": ["理解线性表的逻辑结构"],
            "key_points": ["顺序表与链表"],
            "difficulties": ["指针变化"],
            "teaching_flow": [{"stage": "问题导入", "minutes": 5, "content": "从学生名单引入"}],
            "questions": ["为什么链表随机访问较慢"],
            "exercises": ["实现顺序表插入"],
            "homework": ["比较两种存储结构"],
            "citations": [{"resource_id": "courseware-1", "page": 6, "excerpt": "线性表定义"}],
            "model": "glm-4.5-air",
        }
        with patch(
            "app.api.endpoints.teacher_lesson_prep.lesson_prep_service.summarize",
            new=AsyncMock(return_value=summary_payload),
        ), patch(
            "app.api.endpoints.teacher_lesson_prep.lesson_prep_service.generate_plan",
            new=AsyncMock(return_value=plan_payload),
        ):
            summary_result = asyncio.run(
                summarize_lesson_prep_resources(
                    LessonPrepSummaryRequest(resource_ids=["courseware-1"], query="线性表"),
                    self.teacher_auth,
                )
            )
            plan_result = asyncio.run(
                generate_lesson_plan(
                    LessonPlanGenerateRequest(
                        topic="线性表",
                        duration_minutes=45,
                        resource_ids=["courseware-1"],
                    ),
                    self.teacher_auth,
                )
            )

        self.assertEqual(summary_result["data"]["model"], "glm-4.5-air")
        self.assertEqual(plan_result["data"]["title"], "线性表教学设计")
        self.assertNotIn("lesson-test-key", str(plan_result))

    def test_summary_without_topic_falls_back_to_extracted_courseware_text(self):
        ai_client = AsyncMock()
        ai_client.complete.return_value = {
            "content": "课件介绍计算机组成、程序代码层次与硬件结构。",
            "key_points": ["计算机组成", "程序代码层次"],
        }
        service = TeacherLessonPrepService(ai_client=ai_client)
        resource = next(
            item
            for item in service.catalog.scan()
            if item.frontend_url == "/Computer_ Organization/01_Introduction_fang.pdf"
        )

        result = asyncio.run(service.summarize(LessonPrepSummaryRequest(resource_ids=[resource.id])))

        self.assertEqual(result["content"], "课件介绍计算机组成、程序代码层次与硬件结构。")
        self.assertTrue(result["citations"])
        self.assertEqual(result["citations"][0]["resource_id"], resource.id)
        self.assertEqual(result["citations"][0]["course"], "计算机组成原理")
        self.assertGreaterEqual(result["citations"][0]["page"], 1)
        evidence = json.loads(ai_client.complete.await_args.kwargs["user_prompt"])["evidence"]
        self.assertTrue(evidence)

    def test_summary_without_ai_key_returns_configuration_error(self):
        service = TeacherLessonPrepService.__new__(TeacherLessonPrepService)
        service.retriever = MagicMock()
        search_result = MagicMock()
        search_result.model_dump.return_value = {
            "matches": [{"resource_id": "courseware-1", "page": 1, "excerpt": "courseware evidence"}],
            "index_results": [],
        }
        service.retriever.search.return_value = search_result
        service.ai_client = AsyncMock()
        service.ai_client.complete.side_effect = RuntimeError("AI lesson preparation API key is not configured")

        with self.assertRaises(HTTPException) as ctx:
            asyncio.run(service.summarize(LessonPrepSummaryRequest(resource_ids=["courseware-1"], query="线性表")))

        self.assertEqual(ctx.exception.status_code, 503)
        self.assertIn("AI_LESSON_PREP_API_KEY", ctx.exception.detail)

    def test_draft_storage_is_isolated_by_teacher(self):
        payload = LessonPrepDraftRequest(
            title="线性表教学设计",
            topic="线性表",
            duration_minutes=45,
            resource_ids=["courseware-1"],
            content={"objectives": ["理解线性表"]},
        )
        created = save_lesson_prep_draft(payload, self.teacher_auth, self.db)
        draft_id = created["data"]["draft_id"]
        listed = list_lesson_prep_drafts(self.teacher_auth, self.db)
        detail = get_lesson_prep_draft(draft_id, self.teacher_auth, self.db)

        other_auth = f"Bearer {create_access_token('teacher-2', 'teacher')}"
        other_list = list_lesson_prep_drafts(other_auth, self.db)
        with self.assertRaises(HTTPException) as ctx:
            get_lesson_prep_draft(draft_id, other_auth, self.db)

        self.assertEqual(len(listed["data"]["drafts"]), 1)
        self.assertEqual(detail["data"]["title"], "线性表教学设计")
        self.assertEqual(other_list["data"]["drafts"], [])
        self.assertEqual(ctx.exception.status_code, 404)

    def test_search_requires_selected_resources_and_limits_payload(self):
        with self.assertRaises(ValueError):
            LessonPrepSearchRequest(query="线性表", resource_ids=[])
        with self.assertRaises(ValueError):
            LessonPrepSearchRequest(query="x" * 201, resource_ids=["courseware-1"])
        with self.assertRaises(ValueError):
            LessonPlanGenerateRequest(topic="线性表", resource_ids=[f"resource-{index}" for index in range(11)])
        with self.assertRaises(ValueError):
            LessonPrepDraftRequest(
                title="线性表",
                topic="线性表",
                resource_ids=["courseware-1"],
                content={"text": "x" * 200001},
            )

    def test_unknown_resource_id_becomes_unprocessable_request(self):
        payload = LessonPrepSearchRequest(query="线性表", resource_ids=["missing-resource"])
        with self.assertRaises(HTTPException) as ctx:
            asyncio.run(search_lesson_prep_resources(payload, self.teacher_auth))
        self.assertEqual(ctx.exception.status_code, 422)

    def test_ai_client_uses_dedicated_full_endpoint_and_credentials(self):
        FakeAsyncClient.response_payload = {
            "choices": [{"message": {"content": json.dumps({"content": "ok"}, ensure_ascii=False)}}]
        }
        client = LessonPrepAIClient(client_factory=FakeAsyncClient)
        result = asyncio.run(client.complete(system_prompt="system", user_prompt="user"))

        request = FakeAsyncClient.last_request
        self.assertEqual(result["content"], "ok")
        self.assertEqual(request["url"], settings.AI_LESSON_PREP_BASE_URL)
        self.assertTrue(request["url"].endswith("/chat/completions"))
        self.assertEqual(request["json"]["model"], "glm-4.5-air")
        self.assertEqual(request["headers"]["Authorization"], f"Bearer {settings.AI_LESSON_PREP_API_KEY}")
        self.assertEqual(request["client"]["timeout"], settings.AI_LESSON_PREP_TIMEOUT_SECONDS)

    def test_ai_client_rejects_malformed_or_missing_content(self):
        client = LessonPrepAIClient(client_factory=FakeAsyncClient)
        for payload in ({}, {"choices": []}, {"choices": [{"message": {"content": "not-json"}}]}):
            FakeAsyncClient.response_payload = payload
            with self.subTest(payload=payload), self.assertRaises((RuntimeError, json.JSONDecodeError)):
                asyncio.run(client.complete(system_prompt="system", user_prompt="user"))

    def test_plan_validation_rejects_invalid_flow_and_duration(self):
        service = TeacherLessonPrepService.__new__(TeacherLessonPrepService)
        with self.assertRaises(ValueError):
            service._validate_plan_result({"title": "线性表", "teaching_flow": [{"stage": "导入", "minutes": "bad", "content": "x"}]}, 45, "线性表")
        with self.assertRaises(ValueError):
            service._validate_plan_result({"title": "线性表", "teaching_flow": [{"stage": "导入", "minutes": 5, "content": "x"}]}, 45, "线性表")
        with self.assertRaises(ValueError):
            AILessonPlanResponse.model_validate({
                "title": "线性表",
                "objectives": "not-a-list",
                "key_points": [],
                "difficulties": [],
                "teaching_flow": [{"stage": "导入", "minutes": 45, "content": "x"}],
                "questions": [],
                "exercises": [],
                "homework": [],
            })

    def test_generate_plan_retries_once_when_first_ai_response_is_invalid(self):
        invalid_plan = {
            "title": "计算机组成练习课",
            "objectives": ["理解计算机组成基础"],
            "key_points": ["程序代码层次"],
            "difficulties": ["软件与硬件关系"],
            "teaching_flow": [{"stage": "讲解", "minutes": 40, "content": "讲解课件"}],
            "questions": ["计算机为什么不具备自主智能？"],
            "exercises": ["说明高级语言与汇编语言的区别"],
            "homework": ["复习课件"],
        }
        corrected_plan = {
            **invalid_plan,
            "teaching_flow": [
                {"stage": "讲解", "minutes": 30, "content": "讲解课件"},
                {"stage": "分层练习", "minutes": 15, "content": "完成基础与进阶练习"},
            ],
        }
        service = TeacherLessonPrepService.__new__(TeacherLessonPrepService)
        service.retriever = MagicMock()
        search_result = MagicMock()
        search_result.model_dump.return_value = {
            "matches": [{"resource_id": "courseware-1", "page": 1, "excerpt": "courseware evidence"}],
            "index_results": [],
        }
        service.retriever.search.return_value = search_result
        service.ai_client = AsyncMock()
        service.ai_client.complete.side_effect = [invalid_plan, corrected_plan]

        result = asyncio.run(service.generate_plan(LessonPlanGenerateRequest(
            topic="计算机组成与应用课程介绍",
            duration_minutes=45,
            resource_ids=["courseware-1"],
            requirements="重点生成分层课堂练习",
        )))

        self.assertEqual(service.ai_client.complete.await_count, 2)
        self.assertEqual(sum(item["minutes"] for item in result["teaching_flow"]), 45)
        repair_prompt = service.ai_client.complete.await_args_list[1].kwargs["user_prompt"]
        self.assertIn("duration does not match", repair_prompt)

    def test_summary_validation_rejects_empty_content(self):
        service = TeacherLessonPrepService.__new__(TeacherLessonPrepService)
        with self.assertRaises(ValueError):
            service._validate_summary_result({"content": "", "key_points": []})
        with self.assertRaises(ValueError):
            AILessonSummaryResponse.model_validate({"content": "summary", "key_points": "not-a-list"})

    def test_request_models_forbid_extra_fields(self):
        with self.assertRaises(ValueError):
            LessonPrepSearchRequest.model_validate({
                "query": "线性表",
                "resource_ids": ["courseware-1"],
                "unexpected": "x",
            })

    def test_route_rejects_actual_body_over_limit_even_with_small_declared_length(self):
        from app.api.endpoints.teacher_lesson_prep import router

        app = FastAPI()
        app.include_router(router, prefix="/api")
        client = TestClient(app)
        response = client.post(
            "/api/teacher/lesson-prep/search",
            content=json.dumps({
                "query": "线性表",
                "resource_ids": ["courseware-1"],
                "padding": "x" * (257 * 1024),
            }),
            headers={
                "Authorization": self.teacher_auth,
                "Content-Type": "application/json",
                "Content-Length": "1",
            },
        )
        self.assertEqual(response.status_code, 413)


class TeacherLessonPlanExportDocxTest(unittest.TestCase):
    EXPORT_URL = "/api/teacher/lesson-prep/export-docx"
    DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def setUp(self):
        app = FastAPI()
        app.include_router(router, prefix="/api")
        self.client = TestClient(app)
        self.teacher_auth = f"Bearer {create_access_token('teacher-1', 'teacher')}"
        self.student_auth = f"Bearer {create_access_token('student-1', 'student')}"

    @staticmethod
    def _payload(**overrides):
        payload = {
            "title": "线性表教学设计",
            "topic": "线性表",
            "course_name": "数据结构",
            "audience": "计算机类专业一年级",
            "duration_minutes": 45,
            "objectives": ["理解线性表的逻辑结构", "掌握顺序表插入操作"],
            "key_points": ["顺序表与链表的特点"],
            "difficulties": ["链表指针操作"],
            "teaching_flow": [
                {"stage": "问题导入", "minutes": 5, "content": "从学生名单引入线性表"},
                {"stage": "讲授新知", "minutes": 30, "content": "讲解顺序表与链表"},
                {"stage": "课堂小结", "minutes": 10, "content": "总结线性表要点"},
            ],
            "questions": ["为什么链表随机访问较慢"],
            "exercises": ["实现顺序表插入算法"],
            "homework": ["比较两种存储结构"],
            "summary": "课件围绕线性表的定义与基本操作展开。",
            "citations": [{"name": "线性表课件", "page": 6, "excerpt": "线性表是n个数据元素的有限序列"}],
        }
        payload.update(overrides)
        return payload

    @staticmethod
    def _document_texts(document):
        texts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    texts.append(cell.text)
        return "\n".join(texts)

    def _export(self, payload):
        return self.client.post(self.EXPORT_URL, json=payload, headers={"Authorization": self.teacher_auth})

    def test_export_requires_teacher_token(self):
        response = self.client.post(self.EXPORT_URL, json=self._payload())
        self.assertEqual(response.status_code, 401)

        student_response = self.client.post(self.EXPORT_URL, json=self._payload(), headers={"Authorization": self.student_auth})
        self.assertEqual(student_response.status_code, 403)

    def test_export_returns_valid_docx_with_required_headers(self):
        response = self._export(self._payload())

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"].split(";")[0], self.DOCX_MIME)
        disposition = response.headers["content-disposition"]
        self.assertIn("attachment", disposition)
        self.assertIn("filename*=UTF-8''", disposition)
        self.assertIn(quote("线性表教学设计-教案.docx"), disposition)

        body = response.content
        self.assertTrue(body)
        self.assertTrue(body.startswith(b"PK"))
        self.assertEqual(int(response.headers["content-length"]), len(body))

        document = docx.Document(io.BytesIO(body))
        texts = self._document_texts(document)
        self.assertIn("线性表教学设计", texts)
        self.assertIn("数据结构", texts)
        self.assertIn("教学环节", texts)
        self.assertIn("时长(分钟)", texts)
        self.assertIn("教学内容", texts)
        self.assertIn("从学生名单引入线性表", texts)
        self.assertIn("线性表课件 第6页", texts)
        self.assertEqual(len(document.tables), 2)

    def test_export_rejects_invalid_payload(self):
        missing_title = {key: value for key, value in self._payload().items() if key != "title"}
        invalid_payloads = (
            missing_title,
            self._payload(title="x" * 201),
            self._payload(duration_minutes=0),
            self._payload(duration_minutes=601),
            self._payload(teaching_flow=[{"stage": "导入", "minutes": -1, "content": "内容"}]),
            self._payload(citations=[{"name": "课件", "page": -1}]),
            self._payload(objectives=["知识点"] * 201),
            self._payload(title="线性表", unknown_field="x"),
        )
        for payload in invalid_payloads:
            with self.subTest(payload=str(payload)[:60]):
                response = self._export(payload)
                self.assertEqual(response.status_code, 422)

    def test_export_supports_zero_minutes_stage(self):
        payload = self._payload(teaching_flow=[{"stage": "课前预习检查", "minutes": 0, "content": "快速回顾上节要点"}])
        response = self._export(payload)

        self.assertEqual(response.status_code, 200)
        document = docx.Document(io.BytesIO(response.content))
        texts = self._document_texts(document)
        self.assertIn("课前预习检查", texts)
        self.assertIn("快速回顾上节要点", texts)
        flow_table = document.tables[1]
        flow_cells = [cell.text for row in flow_table.rows for cell in row.cells]
        self.assertIn("0", flow_cells)

    def test_export_skips_empty_sections(self):
        response = self._export(self._payload(questions=[], exercises=[], citations=[]))

        self.assertEqual(response.status_code, 200)
        document = docx.Document(io.BytesIO(response.content))
        texts = self._document_texts(document)
        self.assertNotIn("课堂提问", texts)
        self.assertNotIn("课堂练习", texts)
        self.assertNotIn("课件依据", texts)
        self.assertIn("一、教学目标", texts)
        self.assertIn("四、教学流程", texts)
        self.assertIn("课后作业", texts)

    def test_export_blurbs_and_filenames_are_sanitized(self):
        payload = self._payload(
            title='线性/表:教学*设计?".docx',
            objectives=["  掌握链表插入  ", "   "],
        )
        response = self._export(payload)

        self.assertEqual(response.status_code, 200)
        document = docx.Document(io.BytesIO(response.content))
        texts = self._document_texts(document)
        self.assertIn("1. 掌握链表插入", texts)
        disposition = response.headers["content-disposition"]
        encoded_name = disposition.split("filename*=UTF-8''", 1)[1]
        self.assertEqual(unquote(encoded_name), "线性_表_教学_设计__.docx-教案.docx")

        request = LessonPlanExportRequest.model_validate(payload)
        self.assertEqual(request.objectives, ["掌握链表插入"])
        self.assertIsInstance(request.teaching_flow[0], TeachingStageExport)


if __name__ == "__main__":
    unittest.main()
