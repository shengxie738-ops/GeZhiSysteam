from __future__ import annotations

import importlib.util
import re
from collections import Counter
from pathlib import Path
from typing import Iterable

from app.schemas.teacher_lesson_prep import (
    CoursewareResource,
    CoursewareSearchMatch,
    CoursewareSearchResult,
    DocumentChunk,
    DocumentIndexResult,
)
from app.services.teacher_lesson_prep.courseware_catalog import CoursewareCatalog


_LATIN_WORD = re.compile(r"[a-z0-9]+(?:[-_][a-z0-9]+)*", re.IGNORECASE)
_CJK_RUN = re.compile(r"[\u3400-\u4dbf\u4e00-\u9fff]+")
_WHITESPACE = re.compile(r"\s+")


class DocumentRetriever:
    def __init__(self, catalog: CoursewareCatalog) -> None:
        self.catalog = catalog
        self._index_cache: dict[str, DocumentIndexResult] = {}
        self._index_signatures: dict[str, tuple[int, int]] = {}

    def index_resource(self, resource_id: str, *, refresh: bool = False) -> DocumentIndexResult:
        resource = self.catalog.get_resource(resource_id)
        path = self.catalog.resolve_resource_path(resource_id)
        try:
            stat = path.stat()
        except FileNotFoundError as exc:
            raise ValueError(f"courseware resource is no longer available: {resource_id}") from exc
        signature = (stat.st_mtime_ns, stat.st_size)
        cached = self._index_cache.get(resource_id)
        if cached is not None and self._index_signatures.get(resource_id) == signature and not refresh:
            return cached.model_copy(deep=True)
        if resource.extension == ".ppt":
            result = self._unavailable_result(
                resource,
                "UNSUPPORTED_LEGACY_PPT",
                "legacy .ppt body text indexing is not supported",
            )
        elif resource.extension == ".pdf":
            result = self._index_pdf(resource, path)
        elif resource.extension == ".pptx":
            result = self._index_pptx(resource, path)
        else:
            result = self._unavailable_result(resource, "EXTRACTION_FAILED", "unsupported document type")

        self._index_cache[resource_id] = result
        self._index_signatures[resource_id] = signature
        return result.model_copy(deep=True)

    def search(
        self,
        query: str,
        *,
        resource_ids: Iterable[str] | None = None,
        limit: int = 10,
        fallback_to_chunks: bool = False,
    ) -> CoursewareSearchResult:
        normalized_query = _WHITESPACE.sub(" ", query).strip()
        if not normalized_query:
            return CoursewareSearchResult(query=query)
        if limit < 1:
            raise ValueError("limit must be at least 1")

        selected_ids = list(dict.fromkeys(resource_ids or ()))
        if not selected_ids:
            raise ValueError("at least one courseware resource must be selected")

        chunks: list[DocumentChunk] = []
        chunk_groups: list[list[DocumentChunk]] = []
        index_results: list[DocumentIndexResult] = []
        resources: dict[str, CoursewareResource] = {}
        for resource_id in selected_ids:
            resource = self.catalog.get_resource(resource_id)
            resources[resource_id] = resource
            indexed = self.index_resource(resource_id)
            chunk_groups.append(indexed.chunks)
            index_results.append(indexed.model_copy(update={"chunks": []}))
            chunks.extend(indexed.chunks)

        matches = self.rank_chunks(normalized_query, chunks, limit=limit)
        if not matches and fallback_to_chunks:
            matches = self._fallback_matches(chunk_groups, limit=limit)
        hydrated = []
        for match in matches:
            resource = resources[match.resource_id]
            hydrated.append(
                match.model_copy(
                    update={
                        "course": resource.course,
                        "name": resource.name,
                        "frontend_url": resource.frontend_url,
                    }
                )
            )
        return CoursewareSearchResult(
            query=normalized_query,
            matches=hydrated,
            index_results=index_results,
        )

    def rank_chunks(
        self,
        query: str,
        chunks: Iterable[DocumentChunk],
        *,
        limit: int = 10,
    ) -> list[CoursewareSearchMatch]:
        query_text = _WHITESPACE.sub(" ", query).strip().casefold()
        query_tokens = self._tokenize(query_text)
        if not query_text or not query_tokens or limit < 1:
            return []

        ranked: list[CoursewareSearchMatch] = []
        for chunk in chunks:
            text = _WHITESPACE.sub(" ", chunk.text).strip()
            folded_text = text.casefold()
            text_counts = Counter(self._tokenize(folded_text))
            token_score = sum(min(text_counts[token], 3) for token in query_tokens)
            phrase_score = 4 if query_text in folded_text else 0
            score = float(token_score + phrase_score)
            if score <= 0:
                continue
            ranked.append(
                CoursewareSearchMatch(
                    resource_id=chunk.resource_id,
                    page=chunk.page,
                    excerpt=self._excerpt(text, query_text, query_tokens),
                    score=score,
                )
            )

        ranked.sort(key=lambda match: (-match.score, match.resource_id, match.page))
        return ranked[:limit]

    @staticmethod
    def _fallback_matches(
        chunk_groups: Iterable[list[DocumentChunk]],
        *,
        limit: int,
    ) -> list[CoursewareSearchMatch]:
        groups = [chunks for chunks in chunk_groups if chunks]
        matches = []
        chunk_index = 0
        while len(matches) < limit:
            added = False
            for chunks in groups:
                if chunk_index >= len(chunks):
                    continue
                chunk = chunks[chunk_index]
                excerpt = chunk.text[:600].strip()
                if excerpt:
                    matches.append(
                        CoursewareSearchMatch(
                            resource_id=chunk.resource_id,
                            page=chunk.page,
                            excerpt=excerpt,
                            score=0.1,
                        )
                    )
                    added = True
                    if len(matches) >= limit:
                        break
            if not added:
                break
            chunk_index += 1
        return matches

    def _index_pdf(self, resource: CoursewareResource, path: Path) -> DocumentIndexResult:
        if importlib.util.find_spec("pypdf") is None:
            return self._unavailable_result(resource, "PYPDF_UNAVAILABLE", "pypdf is not installed")
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            chunks = self._page_chunks(resource.id, (page.extract_text() or "" for page in reader.pages))
            return self._indexed_result(resource, chunks)
        except Exception:
            return self._unavailable_result(resource, "EXTRACTION_FAILED", "PDF extraction failed")

    def _index_pptx(self, resource: CoursewareResource, path: Path) -> DocumentIndexResult:
        if importlib.util.find_spec("pptx") is None:
            return self._unavailable_result(
                resource,
                "PYTHON_PPTX_UNAVAILABLE",
                "python-pptx is not installed",
            )
        try:
            from pptx import Presentation

            presentation = Presentation(str(path))
            slide_texts = []
            for slide in presentation.slides:
                texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text]
                slide_texts.append("\n".join(texts))
            chunks = self._page_chunks(resource.id, slide_texts)
            return self._indexed_result(resource, chunks)
        except Exception:
            return self._unavailable_result(resource, "EXTRACTION_FAILED", "PPTX extraction failed")

    @staticmethod
    def _page_chunks(resource_id: str, page_texts: Iterable[str]) -> list[DocumentChunk]:
        chunks = []
        for page_number, raw_text in enumerate(page_texts, start=1):
            text = _WHITESPACE.sub(" ", raw_text).strip()
            if text:
                chunks.append(DocumentChunk(resource_id=resource_id, page=page_number, text=text))
        return chunks

    @staticmethod
    def _indexed_result(
        resource: CoursewareResource,
        chunks: list[DocumentChunk],
    ) -> DocumentIndexResult:
        if not chunks:
            return DocumentIndexResult(
                resource_id=resource.id,
                status="EMPTY_TEXT",
                searchable=False,
                message="document contains no extractable text",
            )
        return DocumentIndexResult(
            resource_id=resource.id,
            status="INDEXED",
            searchable=True,
            chunks=chunks,
        )

    @staticmethod
    def _unavailable_result(
        resource: CoursewareResource,
        status: str,
        message: str,
    ) -> DocumentIndexResult:
        return DocumentIndexResult(
            resource_id=resource.id,
            status=status,
            searchable=False,
            message=message,
        )

    @staticmethod
    def _tokenize(text: str) -> list[str]:
        tokens = _LATIN_WORD.findall(text.casefold())
        for run in _CJK_RUN.findall(text):
            tokens.extend(run)
            tokens.extend(run[index : index + 2] for index in range(len(run) - 1))
        return list(dict.fromkeys(tokens))

    @staticmethod
    def _excerpt(text: str, query_text: str, query_tokens: list[str], width: int = 220) -> str:
        folded = text.casefold()
        positions = [folded.find(query_text)]
        positions.extend(folded.find(token) for token in query_tokens)
        position = min((value for value in positions if value >= 0), default=0)
        start = max(0, position - width // 3)
        end = min(len(text), start + width)
        excerpt = text[start:end].strip()
        if start:
            excerpt = f"...{excerpt}"
        if end < len(text):
            excerpt = f"{excerpt}..."
        return excerpt
